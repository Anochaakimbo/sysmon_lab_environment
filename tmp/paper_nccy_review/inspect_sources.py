from pathlib import Path
import json, hashlib, zipfile
from lxml import etree
from docx import Document
from pypdf import PdfReader

ROOT = Path('S:/Jr.Project(Real)/sysmon-lab')
OUT = ROOT / 'tmp/paper_nccy_review'
OUT.mkdir(parents=True, exist_ok=True)
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

for label, path in [('template',Path('C:/Users/bbcor/Downloads/template_NCCY_full_thai.docx')),('draft',ROOT/'docs/paper_draft_NCCY_v2.docx')]:
    d = Document(path)
    lines = []
    for i,el in enumerate(d.element.body):
        tag = etree.QName(el).localname
        if tag == 'p':
            t = ''.join(el.xpath('.//w:t/text()'))
            lines.append(f'P{i:03d} {t}')
        elif tag == 'tbl':
            lines.append(f'T{i:03d}')
            for row in el.xpath('./w:tr'):
                lines.append(' | '.join(''.join(c.xpath('.//w:t/text()')) for c in row.xpath('./w:tc')))
    (OUT/f'{label}_text.txt').write_text('\n'.join(lines),encoding='utf-8')
    with zipfile.ZipFile(path) as z:
        x = etree.fromstring(z.read('word/document.xml'))
        evidence = {'file':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'sections':[etree.tostring(s,encoding='unicode') for s in x.xpath('//w:sectPr',namespaces=NS)],
            'parts':{n:{'bytes':len(z.read(n)),'sha256':hashlib.sha256(z.read(n)).hexdigest()} for n in z.namelist()},
            'paragraphs':[{'id':i,'text':p.text,'style':p.style.name,'pPr':etree.tostring(p._p.pPr,encoding='unicode') if p._p.pPr is not None else '',
                'runs':[{'text':r.text,'rPr':etree.tostring(r._r.rPr,encoding='unicode') if r._r.rPr is not None else ''} for r in p.runs]} for i,p in enumerate(d.paragraphs)],
            'tables':[etree.tostring(t._tbl,encoding='unicode') for t in d.tables],
            'specials': {name:len(x.xpath(xpath,namespaces=NS)) for name,xpath in [('controls','//w:sdt'),('fields','//w:instrText'),('textboxes','//w:txbxContent'),('tracked','//w:ins | //w:del')]}}
        (OUT/f'{label}_evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
        for name in ('word/styles.xml','word/document.xml'):
            (OUT/f'{label}_{Path(name).name}').write_bytes(z.read(name))
    print(label, 'paragraphs',len(d.paragraphs),'tables',len(d.tables),'images',len(d.inline_shapes),'sections',len(d.sections))
pdf = Path('S:/Jr.Project/Paper/1-s2.0-S277291842500027X-main.pdf')
reader = PdfReader(pdf)
(OUT/'baseline_text.txt').write_text('\n\n'.join(f'PAGE {i+1}\n{p.extract_text(extraction_mode="layout")}' for i,p in enumerate(reader.pages)),encoding='utf-8')
print('baseline pages',len(reader.pages))
